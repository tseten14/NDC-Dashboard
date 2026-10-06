#!/usr/bin/env python3
"""Rebuild the Uganda 2021–2025 case-study briefing from the two final CSVs.

Every number in the document is computed from the CSVs or read from the repo's
NDC configuration. The existing docx is the template for styles, header and footer.
"""
import csv
import json
import re
import subprocess
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / 'data/synthetic/uganda-government-2021-2025'
REFERENCE = FOLDER / 'climate-trace-uganda-sources-2021-2025-with-district.csv'
GOVERNMENT = FOLDER / 'uganda_government_synthetic_2021-2025.csv'
DOCX = ROOT / 'docs/demo/Uganda-2021-2025-Series-Comparison-Case-Studies.docx'
NDC_FRONTEND = ROOT / 'frontend/src/data/uganda-ndc-data.ts'
YEARS = [str(y) for y in range(2021, 2026)]
LABEL = 'SYNTHETIC DEMONSTRATION DATA'
LABEL_TAIL = ' – not official government figures.'
FULL_WIDTH = 9892
SECTORS = [('agriculture', 'Agriculture'), ('manufacturing', 'Manufacturing'),
           ('power', 'Electricity generation'), ('transportation', 'Transport'), ('waste', 'Waste')]
WASTE_DISTRICTS = ['Wakiso', 'Kampala', 'Mukono', 'Kasese']
CEMENT = [('Tororo Cement Tororo Cement Plant', 'Tororo Cement (Tororo)'),
          ('Hima Cement Hima Cement Plant', 'Hima Cement (Kasese)'),
          ('Yaobai International Holding Karamoja Moroto Cement Plant', 'Moroto Cement (Moroto)')]
OTHER_FACILITIES = [('Bidco Uganda Limited', 'Bidco Uganda Limited (Jinja)'),
                    ('Modern Laminates Pulp Mill', 'Modern Laminates Pulp Mill')]
AIRPORTS = ('Kisoro Airport', 'Savannah Airstrip')
ACTION_SECTORS = [('t6', 'Waste'), ('t7', 'Industry (IPPU)'), ('t5', 'Transport'),
                  ('t4', 'Energy'), ('t1', 'AFOLU'), ('t8', 'Agriculture')]


# ---------------------------------------------------------------- data

def read(path):
    with path.open(newline='', encoding='utf-8-sig') as stream:
        rows = list(csv.DictReader(stream))
    for row in rows:
        row['value'] = float(row['emissions_tco2e'])
    return rows


def total(rows, keep=lambda r: True):
    sums = defaultdict(float)
    for row in rows:
        if keep(row):
            sums[row['year']] += row['value']
    return [sums[y] for y in YEARS]


def ratio(synthetic, trace):
    return [s / t if t else float('nan') for s, t in zip(synthetic, trace)]


def pct(value):
    return f'{value * 100:.1f}%'


def direction(value):
    change = (value - 1) * 100
    if abs(change) < 0.05:
        return 'the same as Climate TRACE'
    return f'{abs(change):.1f}% {"higher" if change > 0 else "lower"}'


def kt(value):
    return f'{value / 1e3:.1f}'


def mt(value):
    return f'{value:.1f}' if value == int(value) else f'{value:g}'


def join(items):
    items = list(items)
    return items[0] if len(items) == 1 else f'{", ".join(items[:-1])} and {items[-1]}'


def ndc_config():
    script = ("const t=await import('./config/ndcTargets.js');"
              "const c=await import('./config/ndcCockpitCatalog.js');"
              "console.log(JSON.stringify({targets:t.NDC_TARGETS,meta:c.INDICATOR_META,"
              "activities:c.CATALOG_ACTIVITIES.map(a=>a.body)}))")
    out = subprocess.run(['node', '--input-type=module', '-e', script], cwd=ROOT,
                         check=True, capture_output=True, text=True).stdout
    config = json.loads(out)
    text = NDC_FRONTEND.read_text(encoding='utf-8')
    block = text[text.index('id: "t0"'):text.index('id: "t1"')]
    numbers = {
        'pct': r'by ([\d.]+)% below Business-As-Usual',
        'target': r'full conditional target: ([\d.]+) MtCO',
        'uncond_pct': r'Unconditional contribution: ([\d.]+)% below BAU',
        'uncond': r'below BAU → ([\d.]+) MtCO',
        'cond_pct': r'additional ([\d.]+)% reduction',
        'bau': r'BAU 2030 reference: ([\d.]+) MtCO',
        'base': r'Base year: 2015 \(([\d.]+) MtCO',
    }
    config['economy'] = {k: re.search(p, block).group(1) for k, p in numbers.items()}
    return config


# ------------------------------------------------------------ document

def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)


def para(doc, text='', style=None, lead=None, size=None, italic=False):
    p = doc.add_paragraph(style=style)
    if lead:
        p.add_run(lead).bold = True
    run = p.add_run(text)
    run.italic = italic or None
    if size:
        for r in p.runs:
            r.font.size = Pt(size)
    return p


def label(doc):
    p = doc.add_paragraph()
    first = p.add_run(LABEL)
    first.bold = True
    second = p.add_run(LABEL_TAIL)
    for run in (first, second):
        run.font.size = Pt(9)
    return p


def bullet(doc, text, lead=None):
    return para(doc, text, style='List Bullet', lead=lead)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def heading(doc, text, level):
    return doc.add_paragraph(text, style=f'Heading {level}')


def _cell(row, text, width, header):
    tc = OxmlElement('w:tc')
    tc_pr = OxmlElement('w:tcPr')
    tc_w = OxmlElement('w:tcW')
    tc_w.set(qn('w:type'), 'dxa')
    tc_w.set(qn('w:w'), str(width))
    tc_pr.append(tc_w)
    if header:
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'E9EDF0')
        tc_pr.append(shd)
    tc.append(tc_pr)
    p = OxmlElement('w:p')
    p_pr = OxmlElement('w:pPr')
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), '100')
    spacing.set(qn('w:after'), '100')
    p_pr.append(spacing)
    p.append(p_pr)
    r = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')
    bold = OxmlElement('w:b')
    if not header:
        bold.set(qn('w:val'), '0')
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '18')
    r_pr.extend([bold, size])
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.extend([r_pr, t])
    p.append(r)
    tc.append(p)
    row.append(tc)


def table(doc, headers, rows, widths=None):
    if widths is None:
        first = 2952 if len(headers) == 6 else FULL_WIDTH // len(headers) + FULL_WIDTH % len(headers)
        rest = (FULL_WIDTH - first) // (len(headers) - 1)
        widths = [first] + [rest] * (len(headers) - 1)
    assert len(widths) == len(headers) and all(len(r) == len(headers) for r in rows)
    tbl = OxmlElement('w:tbl')
    tbl_pr = OxmlElement('w:tblPr')
    tbl_w = OxmlElement('w:tblW')
    tbl_w.set(qn('w:type'), 'auto')
    tbl_w.set(qn('w:w'), '0')
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    look = OxmlElement('w:tblLook')
    for k, v in (('firstColumn', '1'), ('firstRow', '1'), ('lastColumn', '0'), ('lastRow', '0'),
                 ('noHBand', '0'), ('noVBand', '1'), ('val', '04A0')):
        look.set(qn(f'w:{k}'), v)
    tbl_pr.extend([tbl_w, layout, look])
    grid = OxmlElement('w:tblGrid')
    for w in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(w))
        grid.append(col)
    tbl.extend([tbl_pr, grid])
    for index, values in enumerate([headers, *rows]):
        tr = OxmlElement('w:tr')
        tr_pr = OxmlElement('w:trPr')
        if index == 0:
            tr_pr.append(OxmlElement('w:tblHeader'))
        tr_pr.append(OxmlElement('w:cantSplit'))
        tr.append(tr_pr)
        for value, width in zip(values, widths):
            _cell(tr, str(value), width, index == 0)
        tbl.append(tr)
    doc.element.body.insert(len(doc.element.body) - 1, tbl)


def set_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    for run in header.runs[1:]:
        run._r.getparent().remove(run._r)
    header.runs[0].text = f'UGANDA DATA COMPARISON  |  {LABEL} – NOT OFFICIAL GOVERNMENT FIGURES'
    footer = section.footer.paragraphs[0]
    footer.runs[0].text = 'Synthetic demonstration data • 2021–2025 • '


# ------------------------------------------------------------- content

def build(trace, synth, config):
    doc = Document(DOCX)
    clear_body(doc)
    set_header_footer(doc)
    matched = lambda r: r['coverage_status'] == 'matched'
    synthetic_only = [r for r in synth if r['coverage_status'] == 'synthetic_only']
    n_synthetic_only = len({r['source_id'] for r in synthetic_only})
    airports_2025 = sum(r['value'] for r in trace if r['source_name'] in AIRPORTS and r['year'] == '2025')
    synth_only_2025 = sum(r['value'] for r in synthetic_only if r['year'] == '2025')
    factors = [float(r['source_variation_factor']) for r in synth
               if r['source_variation_factor'] and float(r['source_variation_factor']) != 1]

    # Cover
    title = doc.add_paragraph(style='Title')
    run = title.add_run('Uganda: comparing')
    run.add_break()
    run.add_text('two datasets')
    para(doc, 'Five proposed Qlik comparisons, with targets and actions | 2021–2025')
    label(doc)
    para(doc, 'This demonstration compares Climate TRACE with a synthetic dataset at national, sector, '
              'district and facility level. It sets those comparisons against Uganda’s NDC targets and '
              'suggests actions that could close the differences it finds.')
    heading(doc, 'What the numbers represent', 2)
    para(doc, 'Climate TRACE is the reference dataset. Its 2021 values retain the supplied records; its '
              '2022–2025 values were created for this demonstration. All values in the synthetic dataset '
              'are illustrative calculations, not official government measurements or targets.')
    para(doc, 'The synthetic national totals were first spread across districts and facilities using Climate '
              f'TRACE shares. Each source was then multiplied by its own factor ({min(factors) * 100:.0f}%–'
              f'{max(factors) * 100:.0f}% of Climate TRACE, taken from the earlier 2021 file and reused for '
              'all five years), and every activity and year was rescaled so it still adds up to the same '
              'national number. Districts and facilities therefore have their own differences. Coverage '
              f'also differs on purpose: {n_synthetic_only} clearly named synthetic facilities appear only '
              f'in the synthetic file and {len(AIRPORTS)} airports appear only in Climate TRACE.')
    para(doc, 'The possible explanations in this briefing describe what to investigate in a real-data '
              'comparison; they are not findings about these facilities.')
    heading(doc, 'Reading the comparison', 2)
    para(doc, 'Climate TRACE is the 100% reference. A synthetic value of 125% means 25% higher; 90% means '
              '10% lower. For example, 125,000 tonnes compared with 100,000 tonnes is 125%.')
    para(doc, 'Amounts use CO₂-equivalent, a common unit for comparing gases that warm the climate. Each '
              'table states its scale. A higher total does not necessarily mean more facilities: the same '
              'facilities may simply have higher estimated values.')
    heading(doc, 'The questions this briefing answers', 2)
    table(doc, ['Question', 'Where', 'What it shows'], [
        ['What are our targets?', 'Uganda’s NDC targets', 'Economy-wide and sector 2030 targets from the Updated NDC 2022'],
        ['Where are the gaps?', 'Comparisons 1–5', 'Where the two datasets differ: total, sector, district, coverage and facility'],
        ['How can we close the gaps?', 'Closing the gaps', 'Data actions for each difference, and the NDC measures for each sector'],
    ], widths=[2700, 2500, 4692])
    para(doc, 'The five comparisons:')
    for text in ('Uganda total: does the difference change over time?',
                 'Sectors: which values are higher or lower?',
                 'Districts: does the same activity compare differently by place?',
                 'Facility coverage: which locations are shared or present in only one file?',
                 'Individual facilities: how do matched cement plants compare each year?'):
        bullet(doc, text)
    heading(doc, 'What the meeting can agree', 2)
    para(doc, 'Confirm the comparison views and which differences the demonstration should show. The final '
              'Qlik dashboard will show a map of Uganda with districts and facility locations; the district '
              'tile grid in the current mock-ups is a placeholder. The synthetic-data label must stay visible '
              'on every interactive page. These comparisons do not establish gaps against Uganda’s climate '
              'commitments.')
    page_break(doc)

    targets_page(doc, config)
    page_break(doc)

    # 1. National
    t_total, s_total = total(trace), total(synth)
    r_total = ratio(s_total, t_total)
    heading(doc, '1. Uganda total', 1)
    label(doc)
    para(doc, 'Is the synthetic total higher or lower, and does the difference change over time?')
    para(doc, 'Uganda • All sectors • 2021–2025', lead='Selection: ')
    para(doc, 'Two annual lines, the selected-year totals, and the synthetic value as a percentage of '
              'Climate TRACE.', lead='Proposed view: ')
    para(doc, 'Amounts in million tonnes CO₂-equivalent.')
    table(doc, ['Dataset / comparison', *YEARS], [
        ['Climate TRACE', *[f'{v / 1e6:.2f}' for v in t_total]],
        ['Synthetic', *[f'{v / 1e6:.2f}' for v in s_total]],
        ['Synthetic (% of TRACE)', *map(pct, r_total)],
    ])
    trend = 'narrows' if abs(r_total[-1] - 1) < abs(r_total[0] - 1) else 'widens'
    para(doc, f'The synthetic total is {direction(r_total[0])} in 2021 and {direction(r_total[-1])} in '
              f'2025. Both totals rise, while the percentage difference {trend}.', lead='Finding: ')
    heading(doc, 'Two possible explanations', 2)
    bullet(doc, 'one dataset may include activities or locations missing from the other.', lead='Broader coverage: ')
    bullet(doc, 'the same activity may be assigned different fuel use or emissions per unit of fuel.',
           lead='Different estimates: ')
    para(doc, 'Compare source lists first, then calculation assumptions for matching records. In the current '
              f'files, facilities found in only one file account for {kt(synth_only_2025)} thousand tonnes '
              f'(synthetic only) and {kt(airports_2025)} thousand tonnes (Climate TRACE only) in 2025, so most '
              'of the national difference comes from the created national inputs.', lead='Check: ')
    para(doc, 'The national total subtracts negative values representing gases taken out of the air. A '
              'comparison of gases released alone would give a different result. National totals are '
              'unchanged by the source-level variation.', lead='Reading note: ')
    assert all(v > 1 for v in r_total), 'Meeting point assumes the synthetic total stays higher'
    change = 'decreases' if trend == 'narrows' else 'increases'
    para(doc, f'The synthetic total remains higher throughout the period, but the relative difference {change}.',
         lead='Meeting point: ')
    page_break(doc)

    # 2. Sectors
    heading(doc, '2. Sector comparison', 1)
    label(doc)
    para(doc, 'Are synthetic values higher in every sector, or does the direction vary?')
    para(doc, 'Uganda • Compare sectors • 2021–2025', lead='Selection: ')
    para(doc, 'Annual percentage comparisons against the 100% Climate TRACE reference, with an option to '
              'view differences in tonnes.', lead='Proposed view: ')
    para(doc, 'Synthetic values as a percentage of Climate TRACE. Above 100% is higher; below 100% is lower.')
    sector_ratios = {}
    for slug, name in SECTORS:
        keep = lambda r, s=slug: r['sector'] == s
        sector_ratios[name] = ratio(total(synth, keep), total(trace, keep))
    table(doc, ['Sector', *YEARS], [[name, *map(pct, sector_ratios[name])] for _, name in SECTORS])
    high = max(sector_ratios, key=lambda n: sector_ratios[n][-1])
    low = min(sector_ratios, key=lambda n: sector_ratios[n][-1])
    para(doc, f'In 2025, synthetic {high.lower()} values are {pct(sector_ratios[high][-1])} of Climate TRACE; '
              f'{low.lower()} values are {pct(sector_ratios[low][-1])}. The difference is not a uniform '
              'increase across sectors.', lead='Finding: ')
    heading(doc, 'Two possible explanations', 2)
    bullet(doc, 'reporting and estimation approaches may capture activities differently across sectors.',
           lead='Different inputs: ')
    bullet(doc, 'an activity may be placed in one sector in one dataset and another sector in the other.',
           lead='Different grouping: ')
    para(doc, 'Align sector definitions, years and source lists. The current files use matching '
              'classifications; their sector differences are created comparison scenarios.', lead='Check: ')
    para(doc, 'The size and direction of the difference depend on the selected activity. Show both '
              'percentages and tonnes before deciding what to investigate.', lead='Meeting point: ')
    page_break(doc)

    district_page(doc, trace, synth)
    page_break(doc)
    coverage_page(doc, trace, synth, synthetic_only)
    page_break(doc)
    facility_page(doc, trace, synth)
    page_break(doc)
    actions_page(doc, config)

    para(doc, 'Sources: the two CSVs in data/synthetic/uganda-government-2021-2025/ and the NDC targets and '
              'measures in config/ndcTargets.js and config/ndcCockpitCatalog.js. Values are rounded for '
              'presentation. Climate TRACE 2021 retains the supplied reference data; later reference years '
              'and all synthetic values are illustrative.')
    for name in (REFERENCE.name, GOVERNMENT.name):
        para(doc, name, size=8)
    return doc


def targets_page(doc, config):
    e = config['economy']
    targets = config['targets']
    heading(doc, 'Uganda’s NDC targets', 1)
    para(doc, 'What are our targets?')
    para(doc, f'Uganda’s Updated NDC (September 2022) aims to reduce economy-wide emissions by {e["pct"]}% '
              f'below business-as-usual (BAU) by 2030: {e["target"]} million tonnes CO₂-equivalent against a '
              f'BAU of {e["bau"]}. The unconditional part, funded domestically, is {e["uncond_pct"]}% below '
              f'BAU ({e["uncond"]} million tonnes); a further {e["cond_pct"]}% depends on international '
              f'finance, technology and capacity support. The base year is 2015 ({e["base"]} million tonnes).')
    para(doc, 'Emission targets for 2030, million tonnes CO₂-equivalent.')
    rows = [['Economy-wide', e['target'], f'{e["pct"]}% below {e["bau"]}', e['base']]]
    for key, name in (('afolu', 'AFOLU (forestry, land use, agriculture)'), ('energy', 'Energy (excluding transport)'),
                      ('transport', 'Transport'), ('waste', 'Waste'), ('ippu', 'Industrial processes (IPPU)')):
        t = targets[key]
        rows.append([name, mt(t['target']),
                     f'{t["reduction_below_bau_pct"]:g}% below {mt(t["bau_2030"])}', mt(t['baseline'])])
    table(doc, ['Sector', '2030 target', 'Reduction below 2030 BAU', '2015 base year'], rows,
          widths=[3700, 1800, 2592, 1800])
    para(doc, 'Agriculture has no standalone emissions target; its measures count towards AFOLU.')
    heading(doc, 'Other sector targets', 2)
    names = {'t2': 'Forest cover', 't9': 'Wetland cover', 't3': 'Electricity generation capacity',
             't10': 'Population with electricity access', 't8': 'Farmers using sustainable land management'}
    rows = []
    for m in config['meta']:
        unit = 'MW' if m['unit'] == 'MW' else '%'
        fmt = (lambda v: f'{v:,g} MW') if unit == 'MW' else (lambda v: f'{v:g}%')
        rows.append([names[m['target_id']], f'{fmt(m["baseline_value"])} ({m["baseline_year"]})',
                     f'{fmt(m["target_value"])} ({m["target_year"]})', m['unit']])
    table(doc, ['Target', 'Baseline', '2030 target', 'Measure'], rows, widths=[3700, 1900, 1900, 2392])
    para(doc, 'Climate TRACE sectors do not map one-to-one to these NDC sectors. For example, the NDC energy '
              'target excludes transport, and AFOLU combines Climate TRACE’s agriculture and forestry. The '
              'synthetic comparison on the following pages measures differences between two datasets; it is '
              'not progress against these targets.', lead='Reading note: ')


def district_page(doc, trace, synth):
    waste = lambda r: r['sector'] == 'waste'
    by = lambda rows, d: total(rows, lambda r: waste(r) and r['district'] == d)
    ratios = {d: ratio(by(synth, d), by(trace, d)) for d in WASTE_DISTRICTS}
    district_trace = defaultdict(float)
    district_synth = defaultdict(float)
    for r in trace:
        if waste(r) and r['year'] == '2025':
            district_trace[r['district']] += r['value']
    for r in synth:
        if waste(r) and r['year'] == '2025':
            district_synth[r['district']] += r['value']
    spread = {d: district_synth[d] / v for d, v in district_trace.items() if v}
    above = sum(v > 1 for v in spread.values())
    below = sum(v < 1 for v in spread.values())
    lo, hi = min(spread, key=spread.get), max(spread, key=spread.get)
    synth_only_districts = sorted({r['district'] for r in synth
                                   if waste(r) and r['coverage_status'] == 'synthetic_only'})
    first, second = WASTE_DISTRICTS[:2]

    heading(doc, '3. District comparison', 1)
    label(doc)
    para(doc, 'For waste, is the difference the same in every district?')
    para(doc, f'Waste • {", ".join(WASTE_DISTRICTS)} • 2021–2025', lead='Selection: ')
    para(doc, 'A map of Uganda with each district shaded by the synthetic value as a percentage of Climate '
              'TRACE; selecting a district shows its two annual lines and a breakdown by type of waste. The '
              'tile grid in the current mock-ups is a placeholder for this map.', lead='Proposed view: ')
    para(doc, 'Synthetic waste values as a percentage of Climate TRACE.')
    table(doc, ['District', *YEARS], [[d, *map(pct, ratios[d])] for d in WASTE_DISTRICTS])
    lowest = min(WASTE_DISTRICTS, key=lambda d: ratios[d][-1])
    highest = max(WASTE_DISTRICTS, key=lambda d: ratios[d][-1])
    para(doc, f'In 2025, the synthetic figure is {direction(ratios[first][-1])} in {first} and '
              f'{direction(ratios[second][-1])} in {second}. Of these four districts, {highest} is furthest '
              f'above Climate TRACE and {lowest} is {direction(ratios[lowest][-1])}. Across all '
              f'{len(spread)} districts with waste records, {above} are higher and {below} lower, from '
              f'{pct(spread[lo])} in {lo} to {pct(spread[hi])} in {hi}'
              f'{", whose total includes synthetic-only waste sites" if hi in synth_only_districts else ""}.',
         lead='Finding: ')

    heading(doc, f'{first} by type of waste, 2025', 2)
    rows = []
    for sub in sorted({r['subsector'] for r in trace if waste(r) and r['district'] == first}):
        keep = lambda r, s=sub: r['subsector'] == s and r['district'] == first and r['year'] == '2025'
        t = sum(r['value'] for r in trace if keep(r))
        s = sum(r['value'] for r in synth if keep(r))
        extra = sum(r['value'] for r in synth if keep(r) and r['coverage_status'] == 'synthetic_only')
        if not t:
            continue
        name = sub.replace('-', ' ').capitalize()
        note = f' (incl. {kt(extra)} synthetic only)' if extra else ''
        rows.append([name, kt(t), f'{kt(s)}{note}', pct(s / t)])
    para(doc, 'Amounts in thousand tonnes CO₂-equivalent.')
    table(doc, ['Type of waste', 'TRACE', 'Synthetic', '% of TRACE'], rows, widths=[3700, 1800, 2592, 1800])

    heading(doc, 'Two possible explanations', 2)
    bullet(doc, 'each source can have its own activity or emissions estimate, so districts with different '
                'sources compare differently even within one activity.', lead='Different local estimates: ')
    bullet(doc, f'sites present in only one file change a district’s total. Synthetic-only waste sites sit in '
                f'{join(synth_only_districts)}.', lead='Different coverage: ')
    para(doc, 'Compare waste types separately, then source lists, locations and district boundaries.',
         lead='Check: ')
    para(doc, 'District figures are engineered for the demonstration: national totals spread by Climate TRACE '
              'shares, then varied by a factor for each source and rescaled to the same national number. '
              'They are not independent local measurements.', lead='Reading note: ')
    para(doc, 'Break a district total into its activities and sources before treating the difference as a '
              'local issue.', lead='Meeting point: ')


def coverage_page(doc, trace, synth, synthetic_only):
    def locations(rows):
        points = [r for r in rows if r['source_type'] == 'point-source' and r['year'] == '2025']
        return {(r['source_name'], r['latitude'], r['longitude']) for r in points}, \
               {r['source_id'] for r in points}
    t_loc, t_ids = locations(trace)
    s_loc, s_ids = locations(synth)
    shared = t_loc & s_loc
    heading(doc, '4. Facility coverage', 1)
    label(doc)
    para(doc, 'Do the files contain the same locations, and which are present in only one file?')
    para(doc, 'Facility comparison • Select year, district or activity', lead='Selection: ')
    para(doc, 'A map of Uganda with facility locations marked as shared, Climate TRACE only or synthetic '
              'only, with counts and the three lists beside it.', lead='Proposed view: ')
    para(doc, 'Current files, 2025. The same coverage is retained in every year.')
    table(doc, ['Coverage measure', 'Climate TRACE', 'Synthetic'], [
        ['Named facility locations', len(t_loc), len(s_loc)],
        ['Point-source IDs', len(t_ids), len(s_ids)],
        ['Locations in both files', len(shared), len(shared)],
        ['Locations only in this file', len(t_loc - s_loc), len(s_loc - t_loc)],
    ], widths=[4000, 2946, 2946])
    only_trace = sorted(n for n, *_ in t_loc - s_loc)
    para(doc, f'Climate TRACE has {len(t_loc)} named locations and the synthetic file {len(s_loc)}. '
              f'{len(shared)} are shared. {" and ".join(only_trace)} appear only in Climate TRACE; '
              f'{len(s_loc - t_loc)} clearly named synthetic facilities appear only in the synthetic file. '
              'Area-wide records, such as county-level agriculture, are excluded.', lead='Finding: ')

    heading(doc, 'Coverage gap example: food processing in Jinja', 2)
    keep = lambda r: r['district'] == 'Jinja' and r['subsector'] == 'food-beverage-tobacco' and r['year'] == '2025'
    t_by, s_by = defaultdict(float), defaultdict(float)
    for r in trace:
        if keep(r):
            t_by[r['source_name']] += r['value']
    for r in synth:
        if keep(r):
            s_by[r['source_name']] += r['value']
    rows = []
    for name in sorted(set(t_by) | set(s_by), key=lambda n: (n.startswith('Synthetic'), n)):
        status = 'In both files' if name in t_by and name in s_by else (
            'Synthetic only' if name in s_by else 'Climate TRACE only')
        rows.append([name, kt(t_by[name]) if name in t_by else 'Not included',
                     kt(s_by[name]) if name in s_by else 'Not included', status])
    rows.append(['Jinja food processing total', kt(sum(t_by.values())), kt(sum(s_by.values())), ''])
    para(doc, 'Amounts in thousand tonnes CO₂-equivalent, 2025.')
    table(doc, ['Facility', 'TRACE', 'Synthetic', 'Coverage'], rows, widths=[4300, 1600, 1600, 2392])
    extra = sum(v for n, v in s_by.items() if n not in t_by)
    n_extra = sum(n not in t_by for n in s_by)
    para(doc, f'Jinja’s food-processing total is {kt(sum(s_by.values()))} thousand tonnes in the synthetic file '
              f'and {kt(sum(t_by.values()))} in Climate TRACE. {kt(extra)} thousand tonnes come from {n_extra} '
              'synthetic plants that Climate TRACE does not include.', lead='Finding: ')
    heading(doc, 'Two possible explanations', 2)
    bullet(doc, 'one list may include newer, smaller or previously undocumented facilities.',
           lead='Additional coverage: ')
    bullet(doc, 'one list may count a whole site once, while another counts separate operations or duplicates.',
           lead='Different counting: ')
    para(doc, 'Match identifiers, names and coordinates; review dates and inclusion rules; remove duplicates '
              'before counting sites.', lead='Check: ')
    para(doc, 'A larger count does not by itself establish a more complete dataset. A facility absent from a '
              'file is “not included”, not zero emissions. The synthetic Jinja plants are fictional.',
         lead='Meeting point: ')


def facility_page(doc, trace, synth):
    def series(rows, name):
        return total(rows, lambda r: r['source_name'] == name)
    heading(doc, '5. Matched facility values', 1)
    label(doc)
    para(doc, 'For the same facility, how different are the values each year?')
    para(doc, 'Manufacturing • Cement • Tororo, Hima and Moroto plants • 2021–2025', lead='Selection: ')
    para(doc, 'The plants on the facility map; selecting one shows its two annual lines, with amounts and a '
              'percentage comparison.', lead='Proposed view: ')
    para(doc, 'Synthetic values as a percentage of Climate TRACE.')
    ratios = {label_: ratio(series(synth, name), series(trace, name)) for name, label_ in CEMENT}
    cement = lambda r: r['subsector'] == 'cement'
    national = ratio(total(synth, cement), total(trace, cement))
    table(doc, ['Cement plant', *YEARS], [*[[l, *map(pct, ratios[l])] for _, l in CEMENT],
                                          ['All cement', *map(pct, national)]])
    first_low = [l for _, l in CEMENT if ratios[l][0] < 1]
    first_high = [l for _, l in CEMENT if ratios[l][0] > 1]
    text = (f'In 2021, {join(first_high)} are higher than Climate TRACE, while {join(first_low)} is lower, '
            f'at {pct(ratios[first_low[0]][0])}. ' if first_low and first_high else '')
    last = {l: ratios[l][-1] for _, l in CEMENT}
    hi = max(last, key=last.get)
    lo = min(last, key=last.get)
    text += (f'By 2025, the plants range from {pct(last[lo])} at {lo} to {pct(last[hi])} at {hi}; all cement '
             f'together is {pct(national[-1])}.')
    para(doc, text, lead='Finding: ')
    para(doc, 'Amounts in thousand tonnes CO₂-equivalent, 2025.')
    rows = []
    for name, label_ in CEMENT + OTHER_FACILITIES:
        t, s = series(trace, name)[-1], series(synth, name)[-1]
        rows.append([label_, kt(t), kt(s), pct(s / t)])
    table(doc, ['Facility', 'TRACE', 'Synthetic', '% of TRACE'], rows, widths=[4300, 1800, 1800, 1992])
    heading(doc, 'Two possible explanations', 2)
    bullet(doc, 'assumed production, fuel use or operating hours may differ.', lead='Different activity estimates: ')
    bullet(doc, 'one figure may cover more operations at the site, or use different emissions assumptions, such '
                'as the clinker share of cement.', lead='Different scope or factors: ')
    para(doc, 'Confirm the same facility, year, activities and gases before comparing production or fuel '
              'inputs and calculation assumptions.', lead='Check: ')
    para(doc, 'Cement plants are used because they are individual facilities in both files. Cattle and other '
              'agriculture records are reported by county, not by farm, so they suit the district view rather '
              'than a facility comparison. Facility values are engineered for the demonstration, not '
              'independent measurements or evidence of incorrect reporting.', lead='Reading note: ')
    para(doc, 'Different facilities can have different comparison ratios, and the direction can change over '
              'time. Match coverage and scope before interpreting them.', lead='Meeting point: ')


def actions_page(doc, config):
    heading(doc, 'Closing the gaps', 1)
    label(doc)
    para(doc, 'How can we close the gaps?')
    para(doc, 'The comparisons find differences between two datasets. Each type of difference points to a '
              'practical data action. Gaps against the NDC targets themselves need confirmed national figures '
              'by sector and agreed comparison thresholds, which this demonstration does not have.')
    table(doc, ['Difference found', 'Action', 'Comparison'], [
        ['Totals differ', 'Agree one inventory scope and year, then reconcile the national totals', '1'],
        ['Sectors differ', 'Agree a sector mapping between the national inventory and Climate TRACE', '2'],
        ['Districts differ', 'Collect local activity data and agree one district boundary set', '3'],
        ['Facilities missing from one file', 'Confirm each site with the operator and add it to both registers', '4'],
        ['Facility values differ', 'Compare production, fuel use and emission factors with the operator', '5'],
    ], widths=[2952, 5440, 1500])
    heading(doc, 'NDC measures by sector', 2)
    para(doc, 'Measures listed in Uganda’s Updated NDC 2022 for the sectors in these comparisons. The '
              'demonstration does not show their effect.')
    rows = []
    for target_id, sector in ACTION_SECTORS:
        for a in config['activities']:
            if a['targetId'] == target_id:
                rows.append([sector, a['name'], a['responsibleMinistry'].replace('Ministry of ', '')])
    table(doc, ['Sector', 'NDC measure', 'Lead ministry'], rows, widths=[2200, 4392, 3300])
    para(doc, 'The cement plants in comparison 5 fall under the clinker substitution measure; the waste '
              'districts in comparison 3 include Kampala, a named Green Cities town.', lead='Link to the data: ')


def main():
    trace, synth = read(REFERENCE), read(GOVERNMENT)
    assert all(r['data_label'] == 'SYNTHETIC_DEMONSTRATION_NOT_OFFICIAL' for r in synth)
    doc = build(trace, synth, ndc_config())
    doc.save(DOCX)
    print(f'Wrote {DOCX.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
